"""Collection-day partition (item #13 in the deposit checklist).

Adds the tweet timestamp, converted to Western Indonesian Time (UTC+7), and the
collection day of every record, and flags the 24-hour window that holds 88.8% of
the corpus (6,334 of 7,131 records). The day column is the grouping used by the
temporal holdout of Table 10.

Timestamps are read from the adjudication workbook, because the published
gold-standard label file carries labels only. The column is
`created_at_validator`, the platform timestamp in the native
"Sat Aug 30 23:38:17 +0000 2025" format.

Usage:
    python code/day_partition.py
    python code/day_partition.py --timestamps data/adjudication_analysis_results.xlsx
"""
from __future__ import annotations

import argparse
import csv
from collections import Counter
from datetime import datetime, timedelta, timezone

from config import DERIVED_DIR, DATA_DIR, WIB_OFFSET_HOURS, WINDOW_START_WIB

TIMESTAMP_FILES = (
    "adjudication_analysis_results.xlsx",
    "210826 - absa demonstrasi dataset (gpt-validator-human-adjudicator).xlsx",
    "human_effort_annotation_reduction_analysis.xlsx",
)
TIMESTAMP_SHEETS = ("Analysis_Data", "Dataset_Flags", "Sheet1")
TIME_FORMATS = ("%a %b %d %H:%M:%S %z %Y", "%Y-%m-%d %H:%M:%S%z", "%Y-%m-%dT%H:%M:%S%z")
WIB = timezone(timedelta(hours=WIB_OFFSET_HOURS))


def parse_timestamp(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=timezone.utc)
    text = str(value).strip()
    for fmt in TIME_FORMATS:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def read_timestamps(path):
    """{index: datetime} from the first workbook/sheet that carries the column."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in ("Analysis_Data", "Dataset_Flags", "Sheet1") + tuple(wb.sheetnames):
            if sheet not in wb.sheetnames:
                continue
            ws = wb[sheet]
            rows = ws.iter_rows(values_only=True)
            header = [str(h).strip() if h is not None else "" for h in next(rows)]
            if "created_at_validator" not in header:
                continue
            i_ix, i_ts = header.index("index"), header.index("created_at_validator")
            out = {}
            for row in rows:
                if row[i_ix] is None:
                    continue
                ts = parse_timestamp(row[i_ts])
                if ts is not None:
                    out[int(row[i_ix])] = ts
            if out:
                return out, sheet
    finally:
        wb.close()
    return {}, None


def locate_timestamps(explicit=None):
    if explicit:
        return explicit
    for name in TIMESTAMP_FILES:
        candidate = DATA_DIR / name
        if candidate.exists():
            return candidate
    raise FileNotFoundError(
        f"no timestamp source found in {DATA_DIR}; expected one of {TIMESTAMP_FILES}")


def window_start():
    return datetime.strptime(WINDOW_START_WIB, "%Y-%m-%d %H:%M").replace(tzinfo=WIB)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--timestamps", default=None, help="workbook with the created_at column")
    args = ap.parse_args()

    path = locate_timestamps(args.timestamps)
    stamps, sheet = read_timestamps(path)
    if not stamps:
        raise SystemExit(f"no created_at_validator column found in {path}")

    start = window_start()
    end = start + timedelta(hours=24)
    days = Counter()
    rows = []
    for index, ts in sorted(stamps.items()):
        wib = ts.astimezone(WIB)
        day = wib.date().isoformat()
        in_window = start <= wib < end
        days[day] += 1
        rows.append(dict(index=index, created_at_utc=ts.strftime("%Y-%m-%d %H:%M:%S"),
                         created_at_wib=wib.strftime("%Y-%m-%d %H:%M:%S"),
                         day=day, in_24h_window=int(in_window)))

    DERIVED_DIR.mkdir(parents=True, exist_ok=True)
    out = DERIVED_DIR / "day_partition.csv"
    with out.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=["index", "created_at_utc", "created_at_wib",
                                           "day", "in_24h_window"])
        w.writeheader()
        w.writerows(rows)

    in_win = sum(r["in_24h_window"] for r in rows)
    print(f"timestamps read from {path.name} [{sheet}]: {len(stamps)}")
    print(f"collection days: {len(days)}")
    for day, n in sorted(days.items()):
        print(f"   {day}: {n}")
    print(f"inside the 24-hour window from {WINDOW_START_WIB} WIB: {in_win} "
          f"({100 * in_win / len(rows):.1f}% of the corpus)")
    print(f"wrote {out}")
    print("\nSelf-check against Section 3: 6,334 records (88.8%) inside one 24-hour window")


if __name__ == "__main__":
    main()
